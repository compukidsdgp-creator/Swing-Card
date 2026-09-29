"""Build the formatted 21-day tracker workbook (openpyxl)."""
from __future__ import annotations

import io
import json
from datetime import datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from .tracker import HOLD

NAVY = "0B1F3A"
BLUE = "2F4BD8"
TEAL = "0E7C86"
GREEN = "1F9D63"
RED = "D64545"
AMBER = "C9860A"
GREY = "F3F5F8"
MID = "D9DEE7"
WHITE = "FFFFFF"
INK = "171A1F"

F_TITLE = Font(name="Calibri", size=18, bold=True, color=WHITE)
F_SUB = Font(name="Calibri", size=10, color="C9D3E6")
F_HEAD = Font(name="Calibri", size=10, bold=True, color=WHITE)
F_BOLD = Font(name="Calibri", size=10, bold=True, color=INK)
F_BODY = Font(name="Calibri", size=10, color=INK)
F_MUTED = Font(name="Calibri", size=9, italic=True, color="6B7280")
FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_HEAD = PatternFill("solid", fgColor="1E3A66")
FILL_GREY = PatternFill("solid", fgColor=GREY)
FILL_CUR = PatternFill("solid", fgColor="FFF4D6")
FILL_FUT = PatternFill("solid", fgColor="FAFAFB")
THIN = Side(style="thin", color=MID)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")

INR = '₹#,##,##0;[Red]-₹#,##,##0'
INR2 = '#,##0.00'
PCT = '+0.00%;[Red]-0.00%;0.00%'

FIRST_DAY_COL = 6  # column F = Day 0 on the grid sheets
HDR_ROW = 4        # header row on grid sheets
DATE_ROW = 5       # date row
DATA_ROW = 6       # first stock row


def _banner(ws, title: str, sub: str, width: int):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws["A1"].value, ws["A2"].value = title, sub
    ws["A1"].font, ws["A2"].font = F_TITLE, F_SUB
    for r in (1, 2):
        for c in range(1, width + 1):
            ws.cell(r, c).fill = FILL_NAVY
        ws.cell(r, 1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 18
    ws.sheet_view.showGridLines = False


def _hdr(cell, text, fill=FILL_HEAD):
    cell.value = text
    cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, fill, CENTER, BORDER


def _body(cell, value=None, fmt=None, bold=False, align=RIGHT, fill=None):
    if value is not None:
        cell.value = value
    cell.font = F_BOLD if bold else F_BODY
    cell.border = BORDER
    cell.alignment = align
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill


def _fmt_date(d: str) -> str:
    return datetime.strptime(d, "%Y-%m-%d").strftime("%d-%b")


# --------------------------------------------------------------------------- #
def build_workbook(state: dict, res: dict) -> bytes:
    b, k = state["batch"], res["kpi"]
    stocks = b["stocks"]
    n = len(stocks)
    cur = k["day"]
    wb = Workbook()

    last_r = DATA_ROW + n - 1
    tot_r = last_r + 1
    day_col = lambda i: FIRST_DAY_COL + i  # noqa: E731
    last_day_col = day_col(HOLD)
    sub = (f"Signal {b['signal_date']}  •  Day {cur} of {HOLD}  •  {k['phase']}  •  "
           f"Exit ~{k['exit_date']}  •  Updated {datetime.now():%d-%b-%Y %H:%M}")

    # ============================ DAILY CLOSES ============================ #
    wc = wb.active
    wc.title = "Daily Closes"
    width = last_day_col + 5
    _banner(wc, f"SwingScope Batch {b['batch_id']} — Daily Close Tracker", sub, width)
    wc.cell(3, 1, "Closes from Day 0 (signal) to Day 21 (exit). Highlighted column = latest session. "
                  "Grey = future session (projected date).").font = F_MUTED

    fixed = ["#", "Symbol", "ISIN", "Qty", "Entry ₹"]
    for i, h in enumerate(fixed, start=1):
        _hdr(wc.cell(HDR_ROW, i), h)
        wc.merge_cells(start_row=HDR_ROW, start_column=i, end_row=DATE_ROW, end_column=i)
    for i in range(HOLD + 1):
        c = day_col(i)
        fill = PatternFill("solid", fgColor=AMBER) if i == cur else (
            PatternFill("solid", fgColor=TEAL) if i in (0, 1, HOLD) else FILL_HEAD)
        _hdr(wc.cell(HDR_ROW, c), f"Day {i}", fill)
        d = res["all_dates"][i] if i < len(res["all_dates"]) else ""
        dc = wc.cell(DATE_ROW, c, _fmt_date(d) + ("" if i <= cur else " (p)") if d else "")
        dc.font = Font(name="Calibri", size=9, bold=i <= cur, color=WHITE if i <= cur else "475569")
        dc.fill = PatternFill("solid", fgColor="334E7E") if i <= cur else PatternFill("solid", fgColor=MID)
        dc.alignment, dc.border = CENTER, BORDER
    extra = ["Last Close", "Return %", "P&L ₹", "Day Chg %", "Trend"]
    for j, h in enumerate(extra):
        _hdr(wc.cell(HDR_ROW, last_day_col + 1 + j), h, PatternFill("solid", fgColor=NAVY))
        wc.merge_cells(start_row=HDR_ROW, start_column=last_day_col + 1 + j,
                       end_row=DATE_ROW, end_column=last_day_col + 1 + j)

    closes, entry = res["closes"], res["entry"]
    first_dc, last_dc = L(day_col(0)), L(last_day_col)
    for idx, s in enumerate(stocks):
        r = DATA_ROW + idx
        zebra = FILL_GREY if idx % 2 else None
        _body(wc.cell(r, 1), s["rank"], align=CENTER, fill=zebra)
        _body(wc.cell(r, 2), s["symbol"], bold=True, align=LEFT, fill=zebra)
        _body(wc.cell(r, 3), s["isin"], align=LEFT, fill=zebra)
        wc.cell(r, 3).font = Font(name="Calibri", size=9, color="6B7280")
        _body(wc.cell(r, 4), float(res["qty"][s["symbol"]]), "0", fill=zebra)
        _body(wc.cell(r, 5), round(float(entry[s["symbol"]]), 2), INR2, bold=True, fill=zebra)
        for i in range(HOLD + 1):
            v = closes.loc[s["symbol"], f"Day {i}"]
            cell = wc.cell(r, day_col(i))
            fill = FILL_CUR if i == cur else (FILL_FUT if i > cur else zebra)
            _body(cell, None if pd.isna(v) or i > cur else round(float(v), 2), INR2, fill=fill)
        rng = f"{first_dc}{r}:{last_dc}{r}"
        lc = last_day_col + 1
        _body(wc.cell(r, lc), f'=LOOKUP(2,1/({rng}<>""),{rng})', INR2, bold=True)
        _body(wc.cell(r, lc + 1), f"={L(lc)}{r}/E{r}-1", PCT, bold=True)
        _body(wc.cell(r, lc + 2), f"=D{r}*({L(lc)}{r}-E{r})", INR, bold=True)
        cur_c, prev_c = L(day_col(cur)), L(day_col(max(cur - 1, 0)))
        _body(wc.cell(r, lc + 3), f"=IFERROR({cur_c}{r}/{prev_c}{r}-1,0)", PCT)
        _body(wc.cell(r, lc + 4), f'=IF({L(lc + 1)}{r}>0,"▲ Up",IF({L(lc + 1)}{r}<0,"▼ Down","■ Flat"))',
              align=CENTER)

    # totals row
    lc = last_day_col + 1
    _body(wc.cell(tot_r, 2), "PORTFOLIO", bold=True, align=LEFT, fill=PatternFill("solid", fgColor=MID))
    for c in (1, 3):
        _body(wc.cell(tot_r, c), "", fill=PatternFill("solid", fgColor=MID))
    _body(wc.cell(tot_r, 4), f"=SUM(D{DATA_ROW}:D{last_r})", "0", bold=True, fill=PatternFill("solid", fgColor=MID))
    _body(wc.cell(tot_r, 5), f"=SUMPRODUCT(D{DATA_ROW}:D{last_r},E{DATA_ROW}:E{last_r})", INR, bold=True,
          fill=PatternFill("solid", fgColor=MID))
    for i in range(HOLD + 1):
        col = L(day_col(i))
        _body(wc.cell(tot_r, day_col(i)),
              f'=IF(COUNT({col}{DATA_ROW}:{col}{last_r})=0,"",SUMPRODUCT($D${DATA_ROW}:$D${last_r},{col}{DATA_ROW}:{col}{last_r}))',
              INR, bold=True, fill=PatternFill("solid", fgColor=MID))
    _body(wc.cell(tot_r, lc), f"=SUMPRODUCT(D{DATA_ROW}:D{last_r},{L(lc)}{DATA_ROW}:{L(lc)}{last_r})", INR, bold=True,
          fill=PatternFill("solid", fgColor=MID))
    _body(wc.cell(tot_r, lc + 1), f"={L(lc)}{tot_r}/E{tot_r}-1", PCT, bold=True, fill=PatternFill("solid", fgColor=MID))
    _body(wc.cell(tot_r, lc + 2), f"=SUM({L(lc + 2)}{DATA_ROW}:{L(lc + 2)}{last_r})", INR, bold=True,
          fill=PatternFill("solid", fgColor=MID))
    _body(wc.cell(tot_r, lc + 3), "", fill=PatternFill("solid", fgColor=MID))
    _body(wc.cell(tot_r, lc + 4), "", fill=PatternFill("solid", fgColor=MID))

    for col in (lc + 1, lc + 3):
        rg = f"{L(col)}{DATA_ROW}:{L(col)}{tot_r}"
        wc.conditional_formatting.add(rg, CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
        wc.conditional_formatting.add(rg, CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))
    wc.conditional_formatting.add(f"{L(lc + 2)}{DATA_ROW}:{L(lc + 2)}{last_r}",
                                  DataBarRule(start_type="min", end_type="max", color="5B8DEF", showValue=True))
    wc.column_dimensions["A"].width = 4
    wc.column_dimensions["B"].width = 14
    wc.column_dimensions["C"].width = 14
    wc.column_dimensions["D"].width = 6
    wc.column_dimensions["E"].width = 10
    for i in range(HOLD + 1):
        wc.column_dimensions[L(day_col(i))].width = 9.5
    for j, w in enumerate([11, 10, 11, 10, 9]):
        wc.column_dimensions[L(lc + j)].width = w
    wc.row_dimensions[HDR_ROW].height = 20
    wc.freeze_panes = wc.cell(DATA_ROW, FIRST_DAY_COL)

    # ============================ RETURNS % =============================== #
    wr = wb.create_sheet("Returns %")
    _banner(wr, "Cumulative Return vs Entry (%)", "Each cell = close ÷ entry − 1. Colour scale: red → white → green.",
            last_day_col)
    wr.cell(3, 1, "Formulas link to 'Daily Closes' — edit a close there and this sheet updates.").font = F_MUTED
    for i, h in enumerate(fixed, start=1):
        _hdr(wr.cell(HDR_ROW, i), h)
    for i in range(HOLD + 1):
        _hdr(wr.cell(HDR_ROW, day_col(i)), f"Day {i}", PatternFill("solid", fgColor=AMBER) if i == cur else FILL_HEAD)
        wr.cell(DATE_ROW, day_col(i), f"='Daily Closes'!{L(day_col(i))}{DATE_ROW}").font = F_MUTED
        wr.cell(DATE_ROW, day_col(i)).alignment = CENTER
    for idx in range(n):
        r = DATA_ROW + idx
        for c in range(1, 6):
            _body(wr.cell(r, c), f"='Daily Closes'!{L(c)}{r}", INR2 if c == 5 else None,
                  bold=c == 2, align=LEFT if c in (2, 3) else RIGHT)
        for i in range(1, HOLD + 1):
            src = f"'Daily Closes'!{L(day_col(i))}{r}"
            _body(wr.cell(r, day_col(i)), f'=IF({src}="","",{src}/$E{r}-1)', PCT)
        _body(wr.cell(r, day_col(0)), "—", align=CENTER)
    r = tot_r
    _body(wr.cell(r, 2), "PORTFOLIO", bold=True, align=LEFT)
    for i in range(1, HOLD + 1):
        src = f"'Daily Closes'!{L(day_col(i))}{tot_r}"
        _body(wr.cell(r, day_col(i)), f"=IF({src}=\"\",\"\",{src}/'Daily Closes'!$E${tot_r}-1)", PCT, bold=True,
              fill=PatternFill("solid", fgColor=MID))
    wr.conditional_formatting.add(
        f"{L(day_col(1))}{DATA_ROW}:{L(last_day_col)}{tot_r}",
        ColorScaleRule(start_type="num", start_value=-0.15, start_color="F8696B",
                       mid_type="num", mid_value=0, mid_color="FFFFFF",
                       end_type="num", end_value=0.15, end_color="63BE7B"))
    for c, w in zip("ABCDE", [4, 14, 14, 6, 10]):
        wr.column_dimensions[c].width = w
    for i in range(HOLD + 1):
        wr.column_dimensions[L(day_col(i))].width = 8.5
    wr.freeze_panes = wr.cell(DATA_ROW, FIRST_DAY_COL)

    # ============================ P&L ₹ ================================== #
    wp = wb.create_sheet("P&L")
    _banner(wp, "Mark-to-Market P&L (₹)", "Qty × (close − entry). Gross of costs.", last_day_col)
    for i, h in enumerate(fixed, start=1):
        _hdr(wp.cell(HDR_ROW, i), h)
    for i in range(HOLD + 1):
        _hdr(wp.cell(HDR_ROW, day_col(i)), f"Day {i}", PatternFill("solid", fgColor=AMBER) if i == cur else FILL_HEAD)
    for idx in range(n):
        r = DATA_ROW + idx
        for c in range(1, 6):
            _body(wp.cell(r, c), f"='Daily Closes'!{L(c)}{r}", INR2 if c == 5 else None,
                  bold=c == 2, align=LEFT if c in (2, 3) else RIGHT)
        for i in range(1, HOLD + 1):
            src = f"'Daily Closes'!{L(day_col(i))}{r}"
            _body(wp.cell(r, day_col(i)), f'=IF({src}="","",$D{r}*({src}-$E{r}))', INR)
        _body(wp.cell(r, day_col(0)), "—", align=CENTER)
    _body(wp.cell(tot_r, 2), "TOTAL", bold=True, align=LEFT)
    for i in range(1, HOLD + 1):
        col = L(day_col(i))
        _body(wp.cell(tot_r, day_col(i)), f'=IF(COUNT({col}{DATA_ROW}:{col}{last_r})=0,"",SUM({col}{DATA_ROW}:{col}{last_r}))',
              INR, bold=True, fill=PatternFill("solid", fgColor=MID))
    rg = f"{L(day_col(1))}{DATA_ROW}:{L(last_day_col)}{tot_r}"
    wp.conditional_formatting.add(rg, CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN)))
    wp.conditional_formatting.add(rg, CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED)))
    for c, w in zip("ABCDE", [4, 14, 14, 6, 10]):
        wp.column_dimensions[c].width = w
    for i in range(HOLD + 1):
        wp.column_dimensions[L(day_col(i))].width = 9.5
    wp.freeze_panes = wp.cell(DATA_ROW, FIRST_DAY_COL)

    # ============================ PORTFOLIO =============================== #
    wf = wb.create_sheet("Portfolio Path")
    _banner(wf, "Portfolio Path — Day by Day", "Value, P&L, return, drawdown and Nifty 50 comparison", 9)
    heads = ["Day", "Date", "Status", "Portfolio Value ₹", "P&L ₹", "Return %", "Daily %", "Drawdown %", "Nifty 50 %"]
    for i, h in enumerate(heads, start=1):
        _hdr(wf.cell(4, i), h)
    port = res["portfolio"]
    for d in range(HOLD + 1):
        r = 5 + d
        date_s = res["all_dates"][d] if d < len(res["all_dates"]) else ""
        done = d <= cur
        status = "Signal" if d == 0 else ("Entry" if d == 1 else ("Exit" if d == HOLD else
                  ("Exit reminder" if d == 20 else "Hold")))
        if not done:
            status += " (upcoming)"
        fill = FILL_CUR if d == cur else (None if done else FILL_FUT)
        _body(wf.cell(r, 1), d, "0", align=CENTER, fill=fill)
        _body(wf.cell(r, 2), date_s, align=CENTER, fill=fill)
        _body(wf.cell(r, 3), status, align=LEFT, fill=fill)
        if done:
            p = port.iloc[d]
            vals = [p["Value"], p["P&L"], p["Return %"] / 100,
                    None if pd.isna(p["Daily %"]) else p["Daily %"] / 100,
                    None if pd.isna(p["Drawdown %"]) else p["Drawdown %"] / 100,
                    None if pd.isna(p["Nifty %"]) else p["Nifty %"] / 100]
        else:
            vals = [None] * 6
        for j, (v, f) in enumerate(zip(vals, [INR, INR, PCT, PCT, PCT, PCT]), start=4):
            _body(wf.cell(r, j), None if v is None else round(float(v), 6), f, fill=fill)
    wf.conditional_formatting.add("E5:I26", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN)))
    wf.conditional_formatting.add("E5:I26", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED)))
    for c, w in zip("ABCDEFGHI", [6, 12, 20, 18, 14, 11, 10, 12, 12]):
        wf.column_dimensions[c].width = w
    wf.freeze_panes = "A5"

    ch = LineChart()
    ch.title, ch.height, ch.width = "Portfolio return vs Nifty 50 (%)", 8, 18
    ch.y_axis.title, ch.x_axis.title = "Return", "Day"
    ch.y_axis.number_format = "0.0%"
    ch.add_data(Reference(wf, min_col=6, min_row=4, max_row=5 + HOLD), titles_from_data=True)
    ch.add_data(Reference(wf, min_col=9, min_row=4, max_row=5 + HOLD), titles_from_data=True)
    ch.set_categories(Reference(wf, min_col=1, min_row=5, max_row=5 + HOLD))
    ch.series[0].graphicalProperties.line.solidFill = BLUE
    ch.series[0].graphicalProperties.line.width = 28000
    ch.series[1].graphicalProperties.line.solidFill = "9AA2AD"
    ch.series[1].graphicalProperties.line.dashStyle = "dash"
    for sr in ch.series:
        sr.smooth = False
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    wf.add_chart(ch, "K4")

    # ============================ DASHBOARD =============================== #
    wd = wb.create_sheet("Dashboard", 0)
    _banner(wd, f"SwingScope Momentum Batch {b['batch_id']} — 21-Day Dashboard", sub, 12)
    for c in range(1, 13):
        wd.column_dimensions[L(c)].width = 13
    wd.column_dimensions["A"].width = 2

    # progress bar (row 4)
    wd.cell(4, 2, f"PROGRESS  Day {cur}/{HOLD}").font = F_BOLD
    bar_cells = 10
    filled = round(cur / HOLD * bar_cells)
    for i in range(bar_cells):
        c = wd.cell(5, 2 + i)
        c.fill = PatternFill("solid", fgColor=BLUE if i < filled else MID)
        c.border = Border(left=Side(style="thin", color=WHITE), right=Side(style="thin", color=WHITE))
    wd.cell(5, 12, f"{cur / HOLD:.0%}").font = F_BOLD
    wd.row_dimensions[5].height = 10

    dc = "'Daily Closes'"
    tiles = [
        ("CURRENT DAY", f"Day {cur} / {HOLD}", None, k["phase"]),
        ("LAST SESSION", k["last_session"], None, f"Exit ~{k['exit_date']}"),
        ("INVESTED", f"={dc}!E{tot_r}", INR, "Qty × entry price"),
        ("MARKET VALUE", f"={dc}!{L(lc)}{tot_r}", INR, "At latest close"),
        ("GROSS P&L", f"={dc}!{L(lc + 2)}{tot_r}", INR, "Before costs"),
        ("GROSS RETURN", f"={dc}!{L(lc + 1)}{tot_r}", PCT, "vs entry"),
        ("EST. COSTS", f"={dc}!E{tot_r}*{(b.get('cost_pct') or 0.6) / 100}", INR, f"{b.get('cost_pct') or 0.6}% round trip"),
        ("NET P&L", None, INR, "After est. costs"),
        ("NIFTY 50", None if k["nifty_ret"] is None else k["nifty_ret"] / 100, PCT, "Same window"),
        ("ALPHA vs NIFTY", None if k["alpha"] is None else k["alpha"] / 100, PCT, "Gross return − Nifty"),
        ("WINNERS / LOSERS", f"{k['winners']} / {k['losers']}", None, f"of {n} stocks"),
        ("MAX DRAWDOWN", k["max_dd"] / 100, PCT, "Peak-to-trough, portfolio"),
    ]
    # layout: 4 tiles per row, each tile 3 cols wide x 3 rows
    positions = {}
    for t, (label, val, fmt, note) in enumerate(tiles):
        row = 7 + (t // 4) * 4
        col = 2 + (t % 4) * 3
        positions[label] = (row + 1, col)
        wd.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        wd.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        wd.merge_cells(start_row=row + 2, start_column=col, end_row=row + 2, end_column=col + 1)
        a, v, nn = wd.cell(row, col), wd.cell(row + 1, col), wd.cell(row + 2, col)
        a.value, nn.value = label, note
        v.value = val
        a.font = Font(name="Calibri", size=9, bold=True, color="64748B")
        v.font = Font(name="Calibri", size=16, bold=True, color=NAVY)
        nn.font = Font(name="Calibri", size=8, italic=True, color="94A3B8")
        if fmt:
            v.number_format = fmt
        for rr in range(row, row + 3):
            for cc in (col, col + 1):
                wd.cell(rr, cc).fill = FILL_GREY
                wd.cell(rr, cc).alignment = Alignment(horizontal="left", vertical="center", indent=1)
        wd.cell(row, col).border = Border(top=Side(style="thick", color=BLUE))
        wd.cell(row, col + 1).border = Border(top=Side(style="thick", color=BLUE))
        wd.row_dimensions[row + 1].height = 26
    gr, gc = positions["GROSS P&L"]
    er, ec = positions["EST. COSTS"]
    nr, nc = positions["NET P&L"]
    wd.cell(nr, nc).value = f"={L(gc)}{gr}-{L(ec)}{er}"
    for lbl in ("GROSS P&L", "GROSS RETURN", "NET P&L", "NIFTY 50", "ALPHA vs NIFTY", "MAX DRAWDOWN"):
        r_, c_ = positions[lbl]
        ref = f"{L(c_)}{r_}"
        wd.conditional_formatting.add(ref, CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True, size=16)))
        wd.conditional_formatting.add(ref, CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True, size=16)))

    # stock league table
    top = 20
    wd.cell(top - 1, 2, "STOCK LEAGUE TABLE (sorted by return)").font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    lt_heads = ["Symbol", "Qty", "Entry ₹", "Last ₹", "Return %", "P&L ₹", "Max gain %", "Max loss %", "Weight %", "Contrib %"]
    for j, h in enumerate(lt_heads):
        _hdr(wd.cell(top, 2 + j), h)
    tbl = res["table"].sort_values("Return %", ascending=False).reset_index(drop=True)
    for i, row in tbl.iterrows():
        r = top + 1 + i
        zebra = FILL_GREY if i % 2 else None
        vals = [row["Symbol"], row["Qty"], row["Entry"], row["Last close"], row["Return %"] / 100, row["P&L"],
                row["Max gain %"] / 100, row["Max loss %"] / 100, row["Weight %"] / 100, row["Contribution %"] / 100]
        fmts = [None, "0", INR2, INR2, PCT, INR, PCT, PCT, "0.0%", PCT]
        for j, (v, f) in enumerate(zip(vals, fmts)):
            vv = v if isinstance(v, str) else (None if pd.isna(v) else round(float(v), 6))
            _body(wd.cell(r, 2 + j), vv, f, bold=j == 0, align=LEFT if j == 0 else RIGHT, fill=zebra)
    lt_last = top + len(tbl)
    for col in ("F", "G", "H", "I", "K"):
        rg = f"{col}{top + 1}:{col}{lt_last}"
        wd.conditional_formatting.add(rg, CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
        wd.conditional_formatting.add(rg, CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))
    wd.conditional_formatting.add(f"G{top + 1}:G{lt_last}",
                                  DataBarRule(start_type="min", end_type="max", color="5B8DEF", showValue=True))

    bc = BarChart()
    bc.type, bc.title, bc.height, bc.width = "bar", "Return by stock", 8.5, 14
    bc.add_data(Reference(wd, min_col=6, min_row=top, max_row=lt_last), titles_from_data=True)
    bc.set_categories(Reference(wd, min_col=2, min_row=top + 1, max_row=lt_last))
    bc.y_axis.number_format = "0%"
    bc.legend = None
    bc.series[0].graphicalProperties.solidFill = BLUE
    bc.dataLabels = DataLabelList()
    bc.dataLabels.showVal = True
    bc.dataLabels.showSerName = False
    bc.dataLabels.showCatName = False
    bc.dataLabels.showLegendKey = False
    bc.dataLabels.numFmt = "+0.0%;-0.0%"
    bc.x_axis.delete = False
    bc.y_axis.delete = False
    bc.x_axis.tickLblPos = "low"
    wd.add_chart(bc, f"B{lt_last + 3}")

    lc2 = LineChart()
    lc2.title, lc2.height, lc2.width = "Portfolio value (₹)", 8.5, 14
    lc2.add_data(Reference(wf, min_col=4, min_row=4, max_row=5 + cur), titles_from_data=True)
    lc2.set_categories(Reference(wf, min_col=1, min_row=5, max_row=5 + cur))
    lc2.legend = None
    lc2.series[0].graphicalProperties.line.solidFill = TEAL
    lc2.series[0].graphicalProperties.line.width = 28000
    lc2.series[0].smooth = False
    lc2.y_axis.number_format = '#,##0'
    lc2.x_axis.title = "Day"
    lc2.x_axis.delete = False
    lc2.y_axis.delete = False
    wd.add_chart(lc2, f"H{lt_last + 3}")
    wd.freeze_panes = "A3"

    # ============================ BATCH INFO ============================== #
    wi = wb.create_sheet("Batch Info")
    _banner(wi, "Batch Details (from SwingScope HTML)", b.get("title", ""), 8)
    info = [
        ("Batch ID", b["batch_id"]), ("Signal date (Day 0)", b["signal_date"]),
        ("Exit (HTML estimate)", b.get("exit_date_est")), ("Exit (projected, NSE calendar)", k["exit_date"]),
        ("Hold", f"{HOLD} sessions — no stop, no target"), ("Universe", b.get("universe")),
        ("Capital committed (at signal close)", b.get("capital")), ("Estimated costs", b.get("est_costs")),
        ("Cost %", f"{b.get('cost_pct')}%"), ("Market regime", f"{b.get('regime')} — {b.get('regime_note')}"),
        ("Stocks with a warning", b.get("warnings")),
    ]
    for i, (a, v) in enumerate(info):
        _body(wi.cell(4 + i, 2), a, bold=True, align=LEFT, fill=FILL_GREY)
        _body(wi.cell(4 + i, 3), v, INR if isinstance(v, (int, float)) else None, align=LEFT)
        wi.merge_cells(start_row=4 + i, start_column=3, end_row=4 + i, end_column=8)
    r0 = 6 + len(info)
    heads = ["#", "Symbol", "ISIN", "Qty", "Signal close", "Amount", "12-1 %", "6-1 %", "3-1 %", "Vol %",
             "Traded daily (cr)", "From 12m high %", "Signal rank", "Data quality", "Watch"]
    for j, h in enumerate(heads):
        _hdr(wi.cell(r0, 2 + j), h)
    for i, s in enumerate(stocks):
        vals = [s["rank"], s["symbol"], s["isin"], s["qty"], s["last_close"], s["amount"], s.get("ret_12_1"),
                s.get("ret_6_1"), s.get("ret_3_1"), s.get("volatility"), s.get("traded_daily_cr"),
                s.get("from_12m_high"), s.get("signal_rank"), s.get("data_quality"), s.get("watch")]
        for j, v in enumerate(vals):
            _body(wi.cell(r0 + 1 + i, 2 + j), v, INR2 if j in (4, 5) else None,
                  align=LEFT if j in (1, 2, 13, 14) else RIGHT)
    r1 = r0 + len(stocks) + 3
    wi.cell(r1, 2, "Evidence behind the rule (backtest, from HTML)").font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    for i, e in enumerate(b.get("evidence", [])):
        c = wi.cell(r1 + 1 + i, 2, e)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        wi.merge_cells(start_row=r1 + 1 + i, start_column=2, end_row=r1 + 1 + i, end_column=16)
        wi.row_dimensions[r1 + 1 + i].height = 60
    ca = state.get("corp_actions") or []
    r2 = r1 + 2 + len(b.get("evidence", []))
    wi.cell(r2, 2, "Corporate actions (splits, bonuses, dividends)").font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    if not ca:
        wi.cell(r2 + 1, 2, "None recorded during this hold.").font = F_MUTED
    else:
        for j, h in enumerate(["Symbol", "Ex-date", "Type", "Ratio / amount", "Status", "Source", "Note"]):
            _hdr(wi.cell(r2 + 1, 2 + j), h)
        for i, ev in enumerate(ca):
            val = ev.get("ratio") if ev["kind"] == "split" else ev.get("amount")
            vals = [ev["symbol"], ev["ex_date"], "Split / bonus" if ev["kind"] == "split" else "Dividend",
                    val, ev.get("status", ""), ev.get("source", ""), ev.get("note", "")]
            for j, v in enumerate(vals):
                _body(wi.cell(r2 + 2 + i, 2 + j), v, align=LEFT)
        wi.cell(r2 + 3 + len(ca), 2, "Splits/bonuses: prices before the ex-date are divided by the ratio and the "
                "quantity multiplied, so returns stay continuous. Dividends are shown for information only.").font = F_MUTED
    wi.column_dimensions["A"].width = 2
    wi.column_dimensions["B"].width = 30
    for c in "CDEFGHIJKLMNOP":
        wi.column_dimensions[c].width = 12
    wi.column_dimensions["C"].width = 14

    # ============================ RESEARCH ================================ #
    brief = state.get("research")
    if brief and brief.get("stocks"):
        from .research import VIEW_STYLE, scorecard
        wrs = wb.create_sheet("Research")
        method = {"notes": "your research notes", "auto": "auto-generated", "ai": "Claude + web search"}.get(
            brief.get("method"), "")
        _banner(wrs, "Stock-by-Stock Research & Swing-Trade Shortlist",
                f"Views from {method} · written {brief.get('generated')}", 9)
        heads = ["Rank", "Symbol", "Company", "View", "Conviction", "Why", "Since entry %", "Call check"]
        for j, h in enumerate(heads):
            _hdr(wrs.cell(4, 1 + j), h)
        rows = scorecard(brief, res)
        for i, r in enumerate(rows):
            rr = 5 + i
            bg, fg = VIEW_STYLE.get(r["tone"], VIEW_STYLE["neutral"])
            vals = [r["rank"], r["symbol"], r["name"], r["view"], r["score"], r["why"],
                    None if r["ret"] is None else r["ret"] / 100, r["verdict"] or "awaiting Day 1"]
            for j, v in enumerate(vals):
                c = wrs.cell(rr, 1 + j)
                _body(c, v, {4: "+0.00;-0.00;0.00", 6: PCT}.get(j), bold=j == 1,
                      align=LEFT if j in (1, 2, 3, 5, 7) else RIGHT)
                if j == 5:
                    c.alignment = Alignment(wrap_text=True, vertical="center")
            vc = wrs.cell(rr, 4)
            vc.fill = PatternFill("solid", fgColor=bg.lstrip("#"))
            vc.font = Font(name="Calibri", size=10, bold=True, color=fg.lstrip("#"))
            wrs.row_dimensions[rr].height = 30
        last = 4 + len(rows)
        wrs.conditional_formatting.add(f"E5:E{last}", DataBarRule(start_type="num", start_value=-2, end_type="num",
                                                                   end_value=2, color="1F9D63", showValue=True))
        wrs.conditional_formatting.add(f"G5:G{last}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
        wrs.conditional_formatting.add(f"G5:G{last}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))
        # detail
        r0 = last + 3
        wrs.cell(r0 - 1, 1, "STOCK BY STOCK — catalysts, risks and sources").font = Font(name="Calibri", size=12, bold=True, color=NAVY)
        for j, h in enumerate(["Rank", "Symbol", "Company", "Type", "Point", "", "Source", "Link"]):
            _hdr(wrs.cell(r0, 1 + j), h)
        rr = r0 + 1
        kcol = {"catalyst": GREEN, "risk": RED, "context": "64748B"}
        for rank, sym in enumerate(brief["shortlist"], start=1):
            stx = brief["stocks"].get(sym)
            if not stx:
                continue
            for p in stx["points"]:
                vals = [rank, sym, stx["name"], p["kind"].title(), p["text"], None, p.get("source", ""), None]
                for j, v in enumerate(vals):
                    _body(wrs.cell(rr, 1 + j), v, align=LEFT if j else RIGHT)
                wrs.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=6)
                wrs.cell(rr, 5).alignment = Alignment(wrap_text=True, vertical="center")
                wrs.cell(rr, 4).font = Font(name="Calibri", size=10, bold=True, color=kcol.get(p["kind"], "64748B"))
                if p.get("url"):
                    lc_ = wrs.cell(rr, 8, "open ↗")
                    lc_.hyperlink = p["url"]
                    lc_.font = Font(name="Calibri", size=10, color=BLUE, underline="single")
                wrs.row_dimensions[rr].height = 30 if len(p["text"]) < 95 else 44
                rr += 1
        for c, w in zip("ABCDEFGH", [6, 13, 22, 20, 12, 58, 15, 16]):
            wrs.column_dimensions[c].width = w
        wrs.column_dimensions["E"].width = 12
        wrs.freeze_panes = "A5"

    # ============================ NEWS ==================================== #
    news = state.get("news")
    if news and news.get("items"):
        wn = wb.create_sheet("News & Sentiment")
        _banner(wn, "Positive News & Sentiment", f"Fetched {news.get('fetched')}", 7)
        heads = ["Symbol", "Published", "Headline", "Source", "Sentiment", "Label", "Link"]
        for j, h in enumerate(heads):
            _hdr(wn.cell(4, 1 + j), h)
        for i, it in enumerate(news["items"]):
            r = 5 + i
            vals = [it["symbol"], it.get("published", ""), it["title"], it.get("source", ""), it["score"], it["label"], it["link"]]
            for j, v in enumerate(vals):
                _body(wn.cell(r, 1 + j), v, "+0.00;-0.00" if j == 4 else None, align=LEFT if j != 4 else RIGHT)
            wn.cell(r, 7).hyperlink = it["link"]
            wn.cell(r, 7).value = "open ↗"
            wn.cell(r, 7).font = Font(name="Calibri", color=BLUE, underline="single")
        if news["items"]:
            wn.conditional_formatting.add(f"E5:E{4 + len(news['items'])}",
                                          ColorScaleRule(start_type="num", start_value=-1, start_color="F8696B",
                                                         mid_type="num", mid_value=0, mid_color="FFFFFF",
                                                         end_type="num", end_value=1, end_color="63BE7B"))
        for c, w in zip("ABCDEFG", [13, 17, 80, 20, 10, 10, 9]):
            wn.column_dimensions[c].width = w
        wn.freeze_panes = "A5"

    # ============================ STATE (hidden) ========================== #
    ws = wb.create_sheet("_state")
    blob = json.dumps(state, default=str)
    for i in range(0, len(blob), 30000):
        ws.cell(1 + i // 30000, 1, blob[i:i + 30000])
    ws.sheet_state = "hidden"

    for sh in wb.worksheets:
        sh.sheet_properties.tabColor = {"Dashboard": BLUE, "Daily Closes": TEAL, "Returns %": GREEN,
                                        "P&L": AMBER, "Portfolio Path": NAVY, "Research": "8B5CF6"}.get(sh.title, "94A3B8")
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    wb.active = 0

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def state_from_workbook(file_bytes: bytes) -> dict:
    """Restore a batch state from a tracker workbook produced by this app."""
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True)
    if "_state" not in wb.sheetnames:
        raise ValueError("This workbook has no embedded tracker state (was it created by this app?).")
    ws = wb["_state"]
    blob = "".join(str(r[0]) for r in ws.iter_rows(values_only=True) if r and r[0])
    return json.loads(blob)
